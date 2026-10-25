import os
import io

from odoo.tests import TransactionCase, tagged

from odoo.addons.ei_legacy_master.models import parsing

SRC = os.path.join(os.path.dirname(__file__), '..', '..', '..', 'migration', 'source', 'EXI_legacy_codes.csv')


def source_bytes():
    with open(SRC, 'rb') as f:
        return f.read()


@tagged('post_install', '-at_install', 'ei')
class TestParsing(TransactionCase):

    def test_real_csv_counts(self):
        res = parsing.parse_csv(source_bytes())
        self.assertEqual(res.errors, [])
        self.assertEqual(len(res.rows), 639)
        by = {}
        for r in res.rows:
            by[r['sheet_key']] = by.get(r['sheet_key'], 0) + 1
        self.assertEqual(by, parsing.EXPECTED_COUNTS)
        self.assertEqual(res.formula_cells, 0)

    def test_display_code_two_decimals(self):
        for raw, shown in ((4.2, '4.20'), (5.1, '5.10'), (16.1, '16.10'), (44.01, '44.01'), (3, '3.00'), ('4.20', '4.20')):
            self.assertEqual(parsing.display_code(raw), shown)

    def test_float_stripped_code_rejected(self):
        data = source_bytes().replace(b',4.20,', b',4.2,')
        res = parsing.parse_csv(data)
        self.assertTrue(any('4.2' in e and 'N.NN' in e for e in res.errors))

    def test_duplicate_and_missing_row_detected(self):
        lines = source_bytes().decode('utf-8-sig').splitlines()
        res = parsing.parse_csv(('\n'.join(lines + [lines[5]])).encode())
        self.assertTrue(any('duplicate code' in e for e in res.errors))
        res = parsing.parse_csv(('\n'.join(lines[:-1])).encode())
        self.assertTrue(any('expected 79 rows' in e or 'expected 58 rows' in e for e in res.errors))

    def test_formula_injection_sanitised(self):
        self.assertEqual(parsing.sanitize_cell('=1+1'), "'=1+1")
        self.assertEqual(parsing.sanitize_cell('@SUM(A1)'), "'@SUM(A1)")
        self.assertEqual(parsing.sanitize_cell('4.20'), '4.20')

    def test_xlsx_roundtrip_preserves_display(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl missing')
        wb = openpyxl.Workbook()
        wb.active.title = 'INDEX'
        ws = wb.create_sheet('PROFIT & LOSS ')
        ws.append(['CODE', 'NAME', 'NICK'])
        ws.append([])
        ws['A3'], ws['B3'], ws['C3'] = 4.2, 'X', None
        ws['A3'].number_format = '0.00'
        ws['B4'] = '=1+1'
        ws['A4'] = 5.1
        buf = io.BytesIO()
        wb.save(buf)
        res = parsing.parse_xlsx(buf.getvalue())
        self.assertEqual([r['code'] for r in res.rows], ['4.20', '5.10'])
        self.assertEqual(res.rows[0]['sheet'], 'PROFIT & LOSS ')
        self.assertEqual(res.formula_cells, 1)


@tagged('post_install', '-at_install', 'ei')
class TestRealWorkbook(TransactionCase):

    def test_real_xlsx_matches_csv(self):
        xlsx = os.path.join(os.path.dirname(SRC), 'CODE_LIST.xlsx')
        if not os.path.exists(xlsx):
            self.skipTest('CODE_LIST.xlsx not in repo')
        with open(xlsx, 'rb') as f:
            res = parsing.parse_xlsx(f.read())
        self.assertEqual(res.errors, [])
        self.assertEqual(res.formula_cells, 0)
        csv_res = parsing.parse_csv(source_bytes())
        self.assertEqual(res.rows, csv_res.rows)
